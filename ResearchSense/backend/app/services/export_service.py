"""Exports people already have to produce by hand.

A research office compiles an annual return per department (papers by
type, per-person output, citations) and faculty keep publication lists for
CVs and promotion files, usually by copying from Google Scholar into a
spreadsheet. Producing those from the record is what makes keeping the
record current worth a researcher's time, and a department that files its
returns from here has a reason to stay.

Every figure comes from the same data the pages show, so a report never
disagrees with the site.
"""

from __future__ import annotations

import csv
import io
import re
from collections import Counter

from app.core.organisation import document_types
from app.repositories import loader

_BIB_TYPE = {
    "journal": "article",
    "conference": "inproceedings",
    "book-chapter": "incollection",
    "book": "book",
}
_BIB_VENUE = {"article": "journal", "inproceedings": "booktitle", "incollection": "booktitle"}


def _papers_of(researcher_id: int) -> list[dict]:
    return sorted(
        (p for p in loader.load("publications")
         if any(a.get("researcher_id") == researcher_id for a in p.get("authors", []))),
        key=lambda p: (-(p.get("publication_year") or 0), p["title"]),
    )


def _bib_escape(text: str) -> str:
    return re.sub(r"([{}])", r"\\\1", text or "")


def _bib_key(p: dict) -> str:
    first = (p.get("authors") or [{}])[0].get("full_name", "anon").split()
    word = re.sub(r"[^a-z]", "", (p.get("title") or "x").lower().split()[0]) or "x"
    return f"{re.sub(r'[^a-z]', '', (first[-1] if first else 'anon').lower())}{p.get('publication_year') or ''}{word}"


def bibtex(researcher_id: int) -> str:
    entries, seen = [], Counter()
    for p in _papers_of(researcher_id):
        kind = _BIB_TYPE.get(p.get("publication_type") or "journal", "misc")
        key = _bib_key(p)
        seen[key] += 1
        if seen[key] > 1:
            key += chr(ord("a") + seen[key] - 1)
        fields = {
            "title": p["title"],
            "author": " and ".join(a["full_name"] for a in p.get("authors", [])),
            "year": str(p.get("publication_year") or ""),
            _BIB_VENUE.get(kind, "howpublished"): p.get("journal_name") or "",
            "doi": p.get("doi") or "",
        }
        body = ",\n".join(f"  {k} = {{{_bib_escape(v)}}}" for k, v in fields.items() if v)
        entries.append(f"@{kind}{{{key},\n{body}\n}}")
    return "\n\n".join(entries) + "\n"


def publications_csv(researcher_id: int) -> str:
    labels = {t.key: t.label for t in document_types()}
    out = io.StringIO()
    w = csv.writer(out)
    w.writerow(["Year", "Title", "Type", "Venue", "Authors", "DOI", "Citations"])
    for p in _papers_of(researcher_id):
        w.writerow([
            p.get("publication_year") or "",
            p["title"],
            labels.get(p.get("publication_type") or "", p.get("publication_type") or ""),
            p.get("journal_name") or "",
            "; ".join(a["full_name"] for a in p.get("authors", [])),
            p.get("doi") or "",
            p.get("citation_count") or 0,
        ])
    return out.getvalue()


def department_report(department: str, year: int) -> bytes:
    """An annual return for one unit, as an Excel workbook with three sheets:
    a summary, output per person, and the list of papers."""
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.utils import get_column_letter

    people = [r for r in loader.load("researchers")
              if not loader.is_extended(r) and r.get("department") == department]
    if not people:
        raise LookupError("No such department")
    ids = {r["researcher_id"] for r in people}
    papers = [p for p in loader.load("publications")
              if p.get("publication_year") == year
              and any(a.get("researcher_id") in ids for a in p.get("authors", []))]
    labels = {t.key: t.label for t in document_types()}
    by_type = Counter(labels.get(p.get("publication_type") or "", p.get("publication_type"))
                      for p in papers)

    wb = Workbook()
    bold = Font(bold=True)

    s = wb.active
    s.title = "Summary"
    rows = [
        (f"{department}: research output {year}", None),
        (None, None),
        ("People in the directory", len(people)),
        ("Publications in the year", len(papers)),
        ("Citations to those publications (to date)", sum(p.get("citation_count") or 0 for p in papers)),
        ("With international co-authors", sum(1 for p in papers if p.get("international"))),
        (None, None),
        ("By type", None),
        *[(f"  {k}", v) for k, v in by_type.most_common()],
        (None, None),
        ("Source", "ResearchSense; publications as indexed from OpenAlex and faculty submissions. "
                   "Counts include only papers linked to directory profiles."),
    ]
    for row in rows:
        s.append(list(row))
    s["A1"].font = Font(bold=True, size=13)

    per = wb.create_sheet("Per person")
    per.append(["Name", "Rank", "Campus", f"Publications {year}", f"Citations to {year} papers"])
    for r in sorted(people, key=lambda r: r["full_name"]):
        mine = [p for p in papers if any(a.get("researcher_id") == r["researcher_id"]
                                         for a in p.get("authors", []))]
        per.append([r["full_name"], r.get("academic_rank") or r.get("designation") or "",
                    r.get("campus") or "", len(mine),
                    sum(p.get("citation_count") or 0 for p in mine)])

    lst = wb.create_sheet("Publications")
    lst.append(["Title", "Type", "Venue", "Department authors", "All authors", "DOI", "Citations"])
    for p in sorted(papers, key=lambda p: p["title"]):
        lst.append([
            p["title"],
            labels.get(p.get("publication_type") or "", p.get("publication_type") or ""),
            p.get("journal_name") or "",
            "; ".join(a["full_name"] for a in p["authors"] if a.get("researcher_id") in ids),
            "; ".join(a["full_name"] for a in p["authors"]),
            p.get("doi") or "",
            p.get("citation_count") or 0,
        ])

    for sheet in (per, lst):
        for cell in sheet[1]:
            cell.font = bold
        sheet.freeze_panes = "A2"
    for sheet, widths in ((s, (48, 60)), (per, (32, 28, 18, 16, 20)),
                          (lst, (70, 18, 40, 36, 50, 30, 10))):
        for i, width in enumerate(widths, start=1):
            sheet.column_dimensions[get_column_letter(i)].width = width

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

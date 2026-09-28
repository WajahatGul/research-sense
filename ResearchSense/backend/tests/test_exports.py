"""Exports: a CV list and a department's annual return, from the same data."""

import io

from fastapi.testclient import TestClient
from openpyxl import load_workbook

from app.main import app

client = TestClient(app)


def test_bibtex_lists_every_paper_on_the_profile():
    profile = client.get("/api/researchers/8").json()
    r = client.get("/api/researchers/8/export", params={"format": "bibtex"})
    assert r.status_code == 200
    assert "arif-ur-rahman-publications.bib" in r.headers["content-disposition"]
    assert r.text.count("\n@") + r.text.startswith("@") == len(profile["publications"])
    assert "Database Preservation" in r.text


def test_csv_has_a_row_per_paper():
    profile = client.get("/api/researchers/8").json()
    r = client.get("/api/researchers/8/export", params={"format": "csv"})
    rows = [line for line in r.text.splitlines() if line.strip()]
    assert rows[0].startswith("Year,Title,Type")
    assert len(rows) - 1 == len(profile["publications"])


def test_department_report_matches_the_site():
    r = client.get("/api/departments/Computer Science/report", params={"year": 2023})
    assert r.status_code == 200
    wb = load_workbook(io.BytesIO(r.content))
    assert wb.sheetnames == ["Summary", "Per person", "Publications"]
    people = client.get("/api/researchers", params={"department": "Computer Science",
                                                    "page_size": 1}).json()["total"]
    assert wb["Per person"].max_row - 1 == people
    listed = wb["Publications"].max_row - 1
    summary = {row[0].value: row[1].value for row in wb["Summary"].iter_rows()}
    assert summary["Publications in the year"] == listed
    per_person_total = sum(row[3].value for row in wb["Per person"].iter_rows(min_row=2))
    assert per_person_total >= listed  # co-authored papers count for each author


def test_unknown_department_is_404():
    assert client.get("/api/departments/Nope/report", params={"year": 2023}).status_code == 404
